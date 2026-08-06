"""회칙 파일(PDF·docx) 파싱 — T2, 회의 4번 결정.

**이 모듈의 핵심 계약은 "조용히 실패하지 않는다"** 이다. 빈 문자열을 반환하면 인덱싱이
청크 0개로 '성공'하고, 관리자는 회칙을 등록했다고 믿는데 심사는 회칙 없는 팀으로 돈다.
그래서 실패 경로마다 어떤 예외가 나오는지를 전부 고정한다 — 이 프로젝트에서 반복된
사고 유형이라(intake parse_ok · digest advice) 테스트로 못박아 둔다.
"""
import io
import zipfile

import docx
import pytest
from pypdf import PdfWriter

from app.tools import document_parser as P
from app.tools.document_parser import (
    DocumentParseError, EmptyDocumentError, UnsupportedDocumentError, extract_text,
)

# ── fixture 생성 (외부 파일 의존 없이 여기서 만든다) ──────────────────────


def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.rows[r].cells[c].text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf_bytes(lines: list[str]) -> bytes:
    """텍스트 레이어가 있는 최소 PDF. 한글은 폰트 임베딩이 필요해 ASCII로 만든다 —
    추출 경로 자체를 검증하는 것이 목적이고, 한글 추출은 pypdf가 폰트에서 처리한다."""
    ops = b"BT /F1 12 Tf 50 800 Td "
    for line in lines:
        ops += f"({line}) Tj 0 -20 Td ".encode()
    ops += b"ET"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(ops)).encode() + b" >>stream\n" + ops + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\n"
            b"startxref\n" + str(xref).encode() + b"\n%%EOF")
    return bytes(out)


def _blank_pdf() -> bytes:
    """텍스트 레이어가 없는 PDF — 스캔본과 같은 상황이다."""
    w = PdfWriter()
    w.add_blank_page(width=595, height=842)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


RULE_LINES = [
    "제1조 (목적) 이 회칙은 모임 활동비 집행 기준을 정한다.",
    "제2조 (회식비 한도) 1인당 회식비는 3만원을 초과할 수 없다.",
    "제3조 (금지 항목) 개인 용도 물품 구입은 지출로 인정하지 않는다.",
]


# ── 정상 경로 ─────────────────────────────────────────────────────────────

def test_docx_extracts_paragraphs():
    text = extract_text(_docx_bytes(RULE_LINES), "회칙.docx")
    for line in RULE_LINES:
        assert line in text


def test_docx_extracts_table_cells():
    """표 안의 텍스트도 가져온다 — 회칙에 '한도 표'를 표로 넣는 경우가 있고,
    문단만 읽으면 그 조항이 통째로 빠지는데 에러가 안 나서 조용한 누락이 된다."""
    data = _docx_bytes(RULE_LINES, table=[["항목", "한도"], ["교재", "연 5만원"]])
    text = extract_text(data, "회칙.docx")
    assert "교재" in text and "연 5만원" in text


def test_pdf_extracts_text_layer():
    text = extract_text(_pdf_bytes([
        "Article 1 Purpose. This rule defines the expense standards of the group.",
        "Article 2 Meal limit is 30000 KRW per person per meeting.",
    ]), "r.pdf")
    assert "Article 1" in text and "Article 2" in text


def test_content_wins_over_filename():
    """이름과 내용이 다르면 **내용**을 믿는다 — `회칙.pdf`인데 실제로 docx인 경우가 있다."""
    text = extract_text(_docx_bytes(RULE_LINES), "회칙.pdf")
    assert RULE_LINES[0] in text


def test_normalize_puts_article_at_line_start():
    """청킹(`chunk.py`)이 조항 경계를 찾으려면 `제N조`가 줄 앞에 와야 한다.

    PDF 추출물은 줄바꿈이 레이아웃 단위라 조항이 문장 중간에 붙어 나오는데, 그러면
    ARTICLE_PATTERN이 안 걸려 회칙 전체가 한 덩어리로 인덱싱된다(검색 품질 저하).
    """
    one_line = ("제1조 (목적) 이 회칙은 모임 활동비 집행 기준을 정한다. "
                "제2조 (회식비 한도) 1인당 회식비는 3만원을 초과할 수 없다.")
    text = extract_text(_docx_bytes([one_line]), "r.docx")
    assert "\n제2조" in text, "조항이 문장 중간에 붙어 있으면 청킹이 한 덩어리로 만든다"


# ── 실패 경로 — 전부 예외다 (빈 문자열 반환 금지) ─────────────────────────

def test_scanned_pdf_raises_instead_of_returning_empty():
    """스캔본 PDF는 텍스트 레이어가 없다. **여기가 이 모듈에서 가장 중요한 테스트다.**

    빈 문자열을 돌려주면 인덱싱이 청크 0개로 성공하고, 관리자는 회칙을 등록했다고
    믿는데 심사는 회칙 없는 팀으로 돈다. 반드시 실패로 남겨야 관리자가 알 수 있다.
    """
    with pytest.raises(EmptyDocumentError) as e:
        extract_text(_blank_pdf(), "스캔본.pdf")
    assert "스캔" in str(e.value), "관리자가 원인을 알 수 있는 메시지여야 한다"


def test_empty_file_raises():
    with pytest.raises(EmptyDocumentError):
        extract_text(b"", "회칙.pdf")


def test_oversized_file_raises(monkeypatch):
    monkeypatch.setattr(P, "MAX_FILE_BYTES", 100)
    with pytest.raises(DocumentParseError) as e:
        extract_text(b"%PDF-" + b"0" * 200, "big.pdf")
    assert "너무 큽니다" in str(e.value)


@pytest.mark.parametrize("data,filename,hint", [
    (bytes([0, 1, 2]), "회칙.hwp", "hwp"),
    (bytes([0xD0, 0xCF, 0x11, 0xE0]), "회칙.doc", "docx로 저장"),
    ("제1조 목적".encode(), "회칙.txt", "PDF 또는 Word"),
    (bytes([0xFF, 0xD8, 0xFF]), "회칙.jpg", "PDF 또는 Word"),
    (b"random", None, "PDF 또는 Word"),
])
def test_unsupported_formats_raise_with_actionable_message(data, filename, hint):
    """지원 외 형식은 **무엇을 하면 되는지** 알려주는 메시지로 거절한다."""
    with pytest.raises(UnsupportedDocumentError) as e:
        extract_text(data, filename)
    assert hint in str(e.value)


def test_plain_zip_is_not_treated_as_docx():
    """docx는 zip 컨테이너다 — 아무 zip이나 통과하면 안 된다."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a.txt", "hi")
    with pytest.raises(UnsupportedDocumentError):
        extract_text(buf.getvalue(), "회칙.zip")


def test_corrupted_pdf_raises():
    with pytest.raises(DocumentParseError):
        extract_text(b"%PDF-1.4 broken", "회칙.pdf")


def test_filename_says_pdf_but_content_is_garbage():
    """확장자만 맞고 내용이 아니면 '손상' 쪽으로 안내한다 — 형식 미지원과 원인이 다르다."""
    with pytest.raises(DocumentParseError) as e:
        extract_text(b"hello world", "회칙.pdf")
    assert "손상" in str(e.value)


def test_excessively_long_document_raises_instead_of_truncating(monkeypatch):
    """분량 상한 — 자르지 않고 에러를 낸다.

    뒤를 잘라내면 잘린 조항이 심사에서 조용히 빠지고 관리자는 전체가 반영된 줄 안다.
    이 모듈이 피하려는 바로 그 실패 유형이라, 자르는 대신 실패로 남긴다.

    운영 위험도 있다 — 워커는 1대가 순차 처리라 수백 쪽 문서 한 건이 다른 팀
    인덱싱까지 몇 분간 막는다.
    """
    monkeypatch.setattr(P, "MAX_TEXT_CHARS", 200)
    long_rule = ["제%d조 (조항) 이 조항은 지출 기준을 정한다." % i for i in range(1, 40)]
    with pytest.raises(DocumentParseError) as e:
        extract_text(_docx_bytes(long_rule), "긴회칙.docx")
    msg = str(e.value)
    assert "너무 많습니다" in msg and "줄여" in msg, "무엇을 하면 되는지 알려야 한다"
